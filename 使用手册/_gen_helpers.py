# -*- coding: utf-8 -*-
"""DFR 供应链平台使用手册生成器 —— 样式与通用函数"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = RGBColor(0x0E, 0x7C, 0x53)      # 深绿（品牌主色加深，适合打印）
DARK = RGBColor(0x21, 0x25, 0x2B)        # 正文深灰
GRAY = RGBColor(0x6B, 0x72, 0x80)        # 次要文字
RED = RGBColor(0xC0, 0x39, 0x2B)         # 警告
AMBER = RGBColor(0xB4, 0x6A, 0x00)       # 注意
TIP_BG = "E8F5EE"                        # 提示浅绿底
WARN_BG = "FDF3E3"                       # 注意浅黄底
QA_BG = "EEF2F8"                         # QA 浅蓝底

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
EN_FONT = "Calibri"


def set_run_font(run, font=BODY_FONT, size=11, bold=False, color=DARK, en=EN_FONT):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), font)


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=10.5, color=DARK, align=None, font=BODY_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, font=font, size=size, bold=True, color=color)
        else:
            r = p.add_run(part)
            set_run_font(r, font=font, size=size, bold=bold, color=color)


def add_bottom_border(paragraph, color="0E7C53", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number_footer(section, doc_title):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(doc_title + "　·　第 ")
    set_run_font(r, size=8.5, color=GRAY)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run(); r2._element.append(fld1)
    r3 = p.add_run(); r3._element.append(instr)
    r4 = p.add_run(); r4._element.append(fld2)
    set_run_font(r2, size=8.5, color=GRAY)
    set_run_font(r3, size=8.5, color=GRAY)
    set_run_font(r4, size=8.5, color=GRAY)
    r5 = p.add_run(" 页 / 共 ")
    set_run_font(r5, size=8.5, color=GRAY)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "NUMPAGES"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r6 = p.add_run(); r6._element.append(f1)
    r7 = p.add_run(); r7._element.append(it)
    r8 = p.add_run(); r8._element.append(f2)
    set_run_font(r6, size=8.5, color=GRAY)
    set_run_font(r7, size=8.5, color=GRAY)
    set_run_font(r8, size=8.5, color=GRAY)
    r9 = p.add_run(" 页")
    set_run_font(r9, size=8.5, color=GRAY)


def setup_document(meta):
    doc = Document()
    # 页面
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    add_page_number_footer(sec, meta["title"])

    # Normal 样式
    st = doc.styles["Normal"]
    st.font.name = EN_FONT
    st.font.size = Pt(11)
    st.font.color.rgb = DARK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.28

    # 标题样式（黑体）
    h1 = doc.styles["Heading 1"]
    h1.font.name = EN_FONT
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BRAND
    h1.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)

    h2 = doc.styles["Heading 2"]
    h2.font.name = EN_FONT
    h2.font.size = Pt(13.5)
    h2.font.bold = True
    h2.font.color.rgb = BRAND
    h2.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = doc.styles["Heading 3"]
    h3.font.name = EN_FONT
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = DARK
    h3.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)

    return doc


def add_cover(doc, meta):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["title"])
    set_run_font(r, font=HEAD_FONT, size=26, bold=True, color=BRAND)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("subtitle", ""))
    set_run_font(r, font=HEAD_FONT, size=14, bold=False, color=GRAY)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DFR 区块链 · 供应链系统")
    set_run_font(r, font=HEAD_FONT, size=12, color=DARK)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("https://supplychain.de5.net")
    set_run_font(r, size=11, color=BRAND)
    for _ in range(3):
        doc.add_paragraph()
    # 元信息表
    tbl = doc.add_table(rows=len(meta["meta"]), cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta["meta"].items()):
        c0 = tbl.rows[i].cells[0]; c1 = tbl.rows[i].cells[1]
        c0.width = Cm(3.8); c1.width = Cm(9.5)
        set_cell_text(c0, k, bold=True, size=11, color=DARK)
        set_cell_text(c1, v, size=11)
        shade_cell(c0, "EAF3EE")
    doc.add_page_break()


def add_toc_field(doc):
    p = doc.add_paragraph()
    add_bottom_border(p)
    r = p.add_run("目录")
    set_run_font(r, font=HEAD_FONT, size=14, bold=True, color=BRAND)
    p2 = doc.add_paragraph()
    r = p2.add_run("（在 Word 中打开后：全选 → F9 更新域，即可自动生成带页码的目录。WPS：引用 → 更新目录。）")
    set_run_font(r, size=9, color=GRAY)
    p3 = doc.add_paragraph()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin"); fld1.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "（目录将在 Word 中自动生成，请按提示更新域）"
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end")
    r3 = p3.add_run()
    r3._element.append(fld1)
    r4 = p3.add_run()
    r4._element.append(instr)
    r5 = p3.add_run()
    r5._element.append(fld2)
    r6 = p3.add_run(); r6._element.append(t)
    r7 = p3.add_run()
    r7._element.append(fld3)
    doc.add_page_break()


def add_heading(doc, level, text):
    h = doc.add_heading(text, level=level)
    return h


def add_para(doc, text, size=11, bold=False, color=DARK, indent=False, space_after=4):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(space_after)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True, color=color)
        else:
            r = p.add_run(part)
            set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_steps(doc, items, start=1):
    for i, it in enumerate(items, start=start):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        parts = re.split(r"(\*\*[^*]+\*\*)", it)
        first = True
        for part in parts:
            if not part:
                continue
            if first:
                r = p.add_run(f"{i}. ")
                set_run_font(r, size=11, bold=True, color=BRAND)
                first = False
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                set_run_font(r, size=11, bold=True)
            else:
                r = p.add_run(part)
                set_run_font(r, size=11)


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        parts = re.split(r"(\*\*[^*]+\*\*)", it)
        first = True
        for part in parts:
            if not part:
                continue
            if first:
                r = p.add_run("• ")
                set_run_font(r, size=11, bold=True, color=BRAND)
                first = False
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                set_run_font(r, size=11, bold=True)
            else:
                r = p.add_run(part)
                set_run_font(r, size=11)


def add_table(doc, headers, rows, col_widths=None, caption=None):
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        set_run_font(r, font=HEAD_FONT, size=11, bold=True, color=DARK)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = tbl.rows[0].cells[j]
        set_cell_text(c, h, bold=True, size=10.5, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(c, "0E7C53")
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            set_cell_text(tbl.rows[i].cells[j], str(v), size=10.5)
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in tbl.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_tip(doc, text, kind="tip"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    label, color, bg = ("提示", BRAND, TIP_BG) if kind == "tip" else ("注意", AMBER, WARN_BG)
    r = p.add_run(f"【{label}】")
    set_run_font(r, size=10.5, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=DARK)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), bg)
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "14"); left.set(qn("w:space"), "4")
    left.set(qn("w:color"), ("0E7C53" if kind == "tip" else "B46A00"))
    pBdr.append(left)
    pPr.append(pBdr)


def add_qa(doc, items):
    for i, (q, a) in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"Q{i}：{q}")
        set_run_font(r, font=HEAD_FONT, size=11, bold=True, color=BRAND)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), QA_BG)
        pPr.append(shd)
        if isinstance(a, str):
            a = [a]
        for line in a:
            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Cm(0.3)
            ap.paragraph_format.space_after = Pt(2)
            r = ap.add_run("A：")
            set_run_font(r, size=10.5, bold=True, color=DARK)
            parts = re.split(r"(\*\*[^*]+\*\*)", line)
            for part in parts:
                if not part:
                    continue
                if part.startswith("**") and part.endswith("**"):
                    r2 = ap.add_run(part[2:-2])
                    set_run_font(r2, size=10.5, bold=True)
                else:
                    r2 = ap.add_run(part)
                    set_run_font(r2, size=10.5)


def build_doc(meta, blocks):
    doc = setup_document(meta)
    add_cover(doc, meta)
    add_toc_field(doc)
    for blk in blocks:
        kind = blk[0]
        if kind == "h1":
            add_heading(doc, 1, blk[1])
        elif kind == "h2":
            add_heading(doc, 2, blk[1])
        elif kind == "h3":
            add_heading(doc, 3, blk[1])
        elif kind == "p":
            add_para(doc, blk[1])
        elif kind == "steps":
            add_steps(doc, blk[1])
        elif kind == "bullets":
            add_bullets(doc, blk[1])
        elif kind == "table":
            add_table(doc, blk[1], blk[2], col_widths=blk[3] if len(blk) > 3 else None,
                      caption=blk[4] if len(blk) > 4 else None)
        elif kind == "tip":
            add_tip(doc, blk[1], "tip")
        elif kind == "warn":
            add_tip(doc, blk[1], "warn")
        elif kind == "qa":
            add_qa(doc, blk[1])
        elif kind == "pagebreak":
            doc.add_page_break()
    return doc
